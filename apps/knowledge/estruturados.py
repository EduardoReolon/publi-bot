"""DOCX, PPTX e XLSX em Markdown, sem GPU.

Diferente do PDF, estes formatos DECLARAM a estrutura: o paragrafo sabe que e
"Titulo 1", o slide tem titulo, a planilha tem abas e celulas. Nao ha layout a
adivinhar, entao a leitura local e tao confiavel quanto a do Docling — e roda
sem depender da maquina da placa.

A saida usa titulos `#`, que e o que a divisao em blocos da curadoria entende:
cada secao do documento vira um bloco que a pessoa marca ou nao.
"""

from __future__ import annotations

import io

EXTENSOES_ESTRUTURADAS = (".docx", ".pptx", ".xlsx")

# Uma planilha de custos pode ter milhares de linhas. Acima disto a aba e
# cortada: o que vai para a curadoria precisa caber numa tela, e uma tabela
# gigante vetorizada vira ruido na busca.
MAXIMO_DE_LINHAS_POR_ABA = 300


def converter(nome: str, bruto: bytes) -> str:
    nome = nome.lower()
    if nome.endswith(".docx"):
        return docx_para_markdown(bruto)
    if nome.endswith(".pptx"):
        return pptx_para_markdown(bruto)
    if nome.endswith(".xlsx"):
        return xlsx_para_markdown(bruto)
    raise ValueError(f"formato nao estruturado: {nome!r}")


def _celula(valor) -> str:
    texto = "" if valor is None else str(valor)
    return texto.replace("|", "\\|").replace("\n", " ").strip()


def _tabela(linhas: list[list[str]]) -> str:
    linhas = [linha for linha in linhas if any(c.strip() for c in linha)]
    if not linhas:
        return ""
    largura = max(len(linha) for linha in linhas)
    linhas = [linha + [""] * (largura - len(linha)) for linha in linhas]
    cabecalho, *corpo = linhas
    partes = [
        "| " + " | ".join(cabecalho) + " |",
        "| " + " | ".join("---" for _ in cabecalho) + " |",
    ]
    partes.extend("| " + " | ".join(linha) + " |" for linha in corpo)
    return "\n".join(partes)


def docx_para_markdown(bruto: bytes) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    documento = docx.Document(io.BytesIO(bruto))
    partes: list[str] = []

    # Percorre o corpo NA ORDEM, paragrafos e tabelas intercalados. Ler
    # `document.paragraphs` e depois `document.tables` separaria a tabela do
    # texto que a explica.
    for elemento in documento.element.body.iterchildren():
        marca = elemento.tag.rsplit("}", 1)[-1]
        if marca == "p":
            paragrafo = Paragraph(elemento, documento)
            texto = paragrafo.text.strip()
            if not texto:
                continue
            estilo = (paragrafo.style.name if paragrafo.style is not None else "").lower()
            if estilo == "title":
                partes.append(f"# {texto}")
            elif estilo.startswith(("heading", "titulo", "título")):
                digitos = "".join(c for c in estilo if c.isdigit())
                nivel = min(int(digitos or 1) + 1, 4)
                partes.append(f"{'#' * nivel} {texto}")
            elif "list" in estilo:
                partes.append(f"- {texto}")
            else:
                partes.append(texto)
        elif marca == "tbl":
            tabela = Table(elemento, documento)
            linhas = [[_celula(c.text) for c in linha.cells] for linha in tabela.rows]
            if bloco := _tabela(linhas):
                partes.append(bloco)

    return "\n\n".join(partes)


def pptx_para_markdown(bruto: bytes) -> str:
    from pptx import Presentation

    apresentacao = Presentation(io.BytesIO(bruto))
    partes: list[str] = []

    for numero, slide in enumerate(apresentacao.slides, start=1):
        titulo = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            titulo = slide.shapes.title.text_frame.text.strip()
        partes.append(f"## {titulo or f'Slide {numero}'}")

        for forma in slide.shapes:
            if forma == slide.shapes.title:
                continue
            if getattr(forma, "has_table", False) and forma.has_table:
                linhas = [[_celula(c.text) for c in linha.cells] for linha in forma.table.rows]
                if bloco := _tabela(linhas):
                    partes.append(bloco)
            elif getattr(forma, "has_text_frame", False) and forma.has_text_frame:
                for paragrafo in forma.text_frame.paragraphs:
                    texto = "".join(r.text for r in paragrafo.runs).strip()
                    if texto:
                        partes.append(("  " * paragrafo.level) + f"- {texto}")

        if slide.has_notes_slide:
            notas = slide.notes_slide.notes_text_frame.text.strip()
            if notas:
                partes.append(f"Notas: {notas}")

    return "\n\n".join(partes)


def xlsx_para_markdown(bruto: bytes) -> str:
    import openpyxl

    # `data_only`: o valor calculado da formula, e nao a formula. E o numero
    # que alguem leria na planilha.
    livro = openpyxl.load_workbook(io.BytesIO(bruto), read_only=True, data_only=True)
    partes: list[str] = []

    for aba in livro.worksheets:
        linhas = []
        for linha in aba.iter_rows(values_only=True):
            linhas.append([_celula(v) for v in linha])
            if len(linhas) > MAXIMO_DE_LINHAS_POR_ABA:
                break
        cortada = len(linhas) > MAXIMO_DE_LINHAS_POR_ABA
        linhas = linhas[:MAXIMO_DE_LINHAS_POR_ABA]
        bloco = _tabela(linhas)
        if not bloco:
            continue
        partes.append(f"## {aba.title}")
        partes.append(bloco)
        if cortada:
            partes.append(f"(aba cortada nas primeiras {MAXIMO_DE_LINHAS_POR_ABA} linhas)")

    livro.close()
    return "\n\n".join(partes)
