"""Perfil de fonte, citacao sem URL e as novas entradas do acervo.

O produto nasceu citando artigo cientifico: toda fonte tinha URL, autores e
ano, e toda citacao virava link. Estes testes guardam o que muda quando a
fonte e o manual da empresa, a nota do especialista, a tabela oficial do mes
ou um topico de forum:

* nem toda fonte vira link — e a que nao tem URL nao pode derrubar o texto;
* forum nao sustenta a afirmacao central, por mais que o modelo o aponte;
* fonte vencida sai da busca;
* DOCX, PPTX, XLSX e pagina web entram sem GPU, e a nota do especialista
  entra ja curada.
"""

from __future__ import annotations

import datetime
import hashlib
import io
import json

import pytest
from django.core.files.base import ContentFile
from django.urls import reverse
from django.utils import timezone
from django_tenants.utils import schema_context

from apps.content.rendering import Fonte, MarcadorSemFonte, substituir_marcadores
from apps.knowledge.models import Document, DocumentCategory, SuperChunk
from tests.test_interface import ambiente  # noqa: F401


@pytest.fixture(autouse=True)
def _embedding_falso_em_todo_o_arquivo(embedding_falso):
    """Indexar carregaria o modelo real; aqui o que importa e o perfil."""


@pytest.fixture
def tenant_de_fontes(tenant_factory):
    tenant = tenant_factory("fontes")
    with schema_context(tenant.schema_name):
        yield tenant


# ---------------------------------------------------------------------------
# Citacao: link, atribuicao, interna
# ---------------------------------------------------------------------------
def test_atribuicao_vira_o_nome_sem_link():
    fontes = {1: Fonte(url="", anchor="Manual da planilha, v3", modo="atribuicao")}

    texto = substituir_marcadores("Segundo [[FONTE_1]], a aba de BDI vem pronta.", fontes)

    assert texto == "Segundo Manual da planilha, v3, a aba de BDI vem pronta."


def test_link_sem_url_cai_para_atribuicao_em_vez_de_quebrar():
    """Antes, fonte sem URL nem entrava no mapa, e o marcador dela derrubava o
    texto inteiro com `MarcadorSemFonte`."""
    fontes = {1: Fonte(url="", anchor="Entrevista com o eng. Lima", modo="link")}

    assert "](" not in substituir_marcadores("Isso muda tudo [[FONTE_1]].", fontes)


def test_fonte_interna_some_junto_com_a_preposicao():
    fontes = {1: Fonte(url="", anchor="Custos internos 2026", modo="interna")}

    texto = substituir_marcadores(
        "Segundo [[FONTE_1]] o custo cai. O prazo e curto [[FONTE_1]].", fontes
    )

    assert "Custos internos" not in texto
    assert "egundo" not in texto
    assert texto == "O custo cai. O prazo e curto."


def test_referencias_no_fim_misturam_link_e_nome_e_omitem_interna():
    fontes = {
        1: Fonte(url="https://sinapi.exemplo.gov.br/2026-09", anchor="SINAPI, set/2026"),
        2: Fonte(url="", anchor="Manual da planilha", modo="atribuicao"),
        3: Fonte(url="", anchor="Planilha de custos", modo="interna"),
    }

    texto = substituir_marcadores(
        "O preco subiu [[FONTE_1]]. A aba calcula [[FONTE_2]]. Vale para todos [[FONTE_3]].",
        fontes,
        ao_final=True,
    )

    assert "- [SINAPI, set/2026](https://sinapi.exemplo.gov.br/2026-09)" in texto
    assert "- Manual da planilha" in texto
    assert "Planilha de custos" not in texto


def test_marcador_sem_fonte_continua_recusado():
    with pytest.raises(MarcadorSemFonte):
        substituir_marcadores("Algo [[FONTE_9]].", {1: Fonte(url="https://a.b", anchor="A")})


# ---------------------------------------------------------------------------
# Perfil da categoria no indice e na busca
# ---------------------------------------------------------------------------
def _categoria(natureza: str, **extra) -> DocumentCategory:
    from apps.knowledge.perfis import PERFIS

    perfil = {**PERFIS[natureza], **extra}
    return DocumentCategory.objects.create(
        name=natureza,
        slug=f"{natureza}-{DocumentCategory.objects.count()}",
        source_class=natureza,
        **perfil,
    )


def _documento_curado(categoria, texto: str, **campos) -> Document:
    from apps.knowledge.services import indexar_blocos, marcar_curado

    markdown = f"# Documento\n\n## Parte\n\n{texto}"
    bruto = markdown.encode()
    documento = Document.objects.create(
        category=categoria,
        original_file=ContentFile(bruto, name="doc.md"),
        file_sha256=hashlib.sha256(bruto).hexdigest(),
        markdown_full=markdown,
        extraction_method="texto",
        title=campos.pop("title", "Documento"),
        status=Document.Status.PENDING_CURATION,
        **campos,
    )
    from apps.knowledge.blocos import preparar_blocos

    indexar_blocos(
        document=documento, blocos_marcados={b.ordem for b in preparar_blocos(documento)}
    )
    marcar_curado(document=documento, revisado_por=None)
    return documento


@pytest.mark.django_db
def test_o_trecho_carrega_o_perfil_da_categoria(tenant_de_fontes):
    documento = _documento_curado(
        _categoria("empresa"), "A planilha calcula o BDI sozinha.", source_label="Manual v3"
    )

    trecho = documento.chunks.first()
    assert trecho.citation_mode == "atribuicao"
    assert trecho.source_label == "Manual v3"
    assert trecho.supports_central_idea


@pytest.mark.django_db
def test_confidencial_vence_a_escolha_de_citacao(tenant_de_fontes):
    categoria = _categoria("empresa", confidential=True, citation_mode="link")
    documento = _documento_curado(categoria, "Margem interna de 18 por cento.")

    assert documento.chunks.first().citation_mode == "interna"


@pytest.mark.django_db
def test_validade_conta_da_publicacao_e_fonte_vencida_sai_da_busca(tenant_de_fontes):
    from apps.knowledge.services import recuperar

    categoria = _categoria("normativo", validity_days=31)
    antiga = _documento_curado(
        categoria,
        "Tabela de custos de referencia da construcao civil.",
        published_on=timezone.localdate() - datetime.timedelta(days=60),
    )
    nova = _documento_curado(
        categoria,
        "Tabela de custos de referencia da construcao civil atualizada.",
        published_on=timezone.localdate(),
        file_size_bytes=1,
    )

    assert antiga.vencida and not nova.vencida
    _, trechos = recuperar(
        consulta="tabela de custos da construcao", origem="article", distancia_maxima=2.0
    )
    documentos = {t.chunk.document_id for t in trechos}
    assert nova.pk in documentos
    assert antiga.pk not in documentos


@pytest.mark.django_db
def test_sem_data_de_publicacao_a_validade_conta_da_busca(tenant_de_fontes):
    categoria = _categoria("veiculo", validity_days=10)
    documento = _documento_curado(
        categoria, "Reportagem.", fetched_at=timezone.now() - datetime.timedelta(days=3)
    )

    assert documento.valid_until == timezone.localdate() + datetime.timedelta(days=7)


def test_forum_nao_sustenta_a_ideia_central_mesmo_que_o_modelo_diga():
    from apps.content.services import SemEmbasamentoCentral, interpretar_plano

    plano = json.dumps(
        {
            "ideia_central": "Planilha reduz erro de orcamento.",
            "fontes_da_ideia_central": [2],
            "secoes": [
                {"titulo": "A", "objetivo": "a", "fontes": [1], "sustenta_ideia_central": True},
                {"titulo": "B", "objetivo": "b", "fontes": [2]},
            ],
        }
    )

    with pytest.raises(SemEmbasamentoCentral):
        interpretar_plano(plano, total_de_fontes=2, fontes_para_ideia_central={1})


def test_o_modelo_ve_como_usar_cada_fonte():
    from types import SimpleNamespace

    from apps.content.services import montar_contexto_das_fontes

    forum = SimpleNamespace(
        source_authors="",
        source_year=None,
        content="Alguem perguntou como somar o BDI.",
        supports_central_idea=False,
        citation_mode="interna",
    )

    contexto = montar_contexto_das_fontes([forum])

    assert "apenas contexto" in contexto
    assert "citacao interna" in contexto


# ---------------------------------------------------------------------------
# Artigo com fonte sem URL, ponta a ponta
# ---------------------------------------------------------------------------
@pytest.mark.django_db
def test_artigo_com_fonte_da_empresa_sai_sem_link_e_sem_link_de_saida(
    tenant_com_acervo, conexao, monkeypatch
):
    from apps.content.models import Article, Topic
    from apps.ops.models import GenerationJob
    from apps.ops.orchestrator import criar_job
    from tests.test_flows import ROTEIRO_DO_ARTIGO, ModeloFalso, _rodar_ate_o_fim

    # O unico documento do acervo passa a ser material da empresa.
    categoria = DocumentCategory.objects.get()
    categoria.source_class = "empresa"
    categoria.citation_mode = "atribuicao"
    categoria.save()
    SuperChunk.objects.update(citation_mode="atribuicao", source_label="Equipe tecnica da empresa")

    modelo = ModeloFalso(ROTEIRO_DO_ARTIGO)
    monkeypatch.setattr("apps.content.inference.get_provider", lambda *a, **k: modelo)
    topic = Topic.objects.create(title="Efeito no metabolismo")
    job = criar_job(kind=GenerationJob.Kind.PILLAR_ARTICLE, target_object_id=str(topic.pk))

    assert _rodar_ate_o_fim(str(job.pk)) == GenerationJob.Status.DONE

    artigo = Article.objects.get(topic=topic)
    assert "Equipe tecnica da empresa" in artigo.body_markdown
    assert "revista.exemplo.org" not in artigo.body_markdown
    assert artigo.outbound_link_url == ""
    assert artigo.primary_source is None


# ---------------------------------------------------------------------------
# Documentos estruturados, sem GPU
# ---------------------------------------------------------------------------
def test_docx_vira_markdown_com_titulos_e_tabela_na_ordem():
    import docx

    from apps.knowledge.estruturados import docx_para_markdown

    documento = docx.Document()
    documento.add_heading("Como usar a planilha", level=1)
    documento.add_paragraph("Preencha a aba de insumos primeiro.")
    tabela = documento.add_table(rows=2, cols=2)
    tabela.cell(0, 0).text, tabela.cell(0, 1).text = "Item", "Preco"
    tabela.cell(1, 0).text, tabela.cell(1, 1).text = "Cimento", "35,00"
    documento.add_paragraph("Depois, confira o BDI.")
    memoria = io.BytesIO()
    documento.save(memoria)

    markdown = docx_para_markdown(memoria.getvalue())

    assert "## Como usar a planilha" in markdown
    assert "| Cimento | 35,00 |" in markdown
    assert markdown.index("Cimento") < markdown.index("confira o BDI")


def test_pptx_vira_um_bloco_por_slide():
    from pptx import Presentation

    from apps.knowledge.estruturados import pptx_para_markdown

    apresentacao = Presentation()
    slide = apresentacao.slides.add_slide(apresentacao.slide_layouts[1])
    slide.shapes.title.text = "Por que orcar por etapa"
    slide.placeholders[1].text = "Evita esquecer servicos"
    memoria = io.BytesIO()
    apresentacao.save(memoria)

    markdown = pptx_para_markdown(memoria.getvalue())

    assert "## Por que orcar por etapa" in markdown
    assert "- Evita esquecer servicos" in markdown


def test_xlsx_vira_uma_tabela_por_aba_com_teto_de_linhas():
    import openpyxl

    from apps.knowledge.estruturados import MAXIMO_DE_LINHAS_POR_ABA, xlsx_para_markdown

    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Insumos"
    aba.append(["Item", "Preco"])
    for n in range(MAXIMO_DE_LINHAS_POR_ABA + 10):
        aba.append([f"item {n}", n])
    memoria = io.BytesIO()
    livro.save(memoria)

    markdown = xlsx_para_markdown(memoria.getvalue())

    assert "## Insumos" in markdown
    assert "| Item | Preco |" in markdown
    assert "aba cortada" in markdown


# ---------------------------------------------------------------------------
# Pagina web
# ---------------------------------------------------------------------------
PAGINA = """<html><head><title>Como calcular o BDI</title>
<meta name="author" content="Ana Engenheira">
<meta property="article:published_time" content="2026-03-10">
</head><body><nav>Inicio | Blog | Contato</nav><article><h1>Como calcular o BDI</h1>
<p>O BDI, ou Beneficios e Despesas Indiretas, e o percentual aplicado sobre o custo
direto de uma obra para cobrir despesas que nao aparecem em nenhum servico
especifico, como administracao central, seguros, riscos e o lucro da empresa.</p>
<p>A formula mais usada combina essas parcelas de forma multiplicativa, e nao
somando tudo, porque cada uma incide sobre a base ja acrescida das anteriores.
Somar direto subestima o resultado final do orcamento.</p></article>
<footer>Todos os direitos reservados</footer></body></html>"""


def test_pagina_vira_texto_principal_com_data_e_autor():
    from apps.knowledge.web import extrair_pagina

    pagina = extrair_pagina(PAGINA.encode(), url="https://blog.exemplo.com.br/bdi")

    assert "Despesas Indiretas" in pagina.markdown
    assert "direitos reservados" not in pagina.markdown
    assert pagina.titulo == "Como calcular o BDI"
    assert pagina.data == datetime.date(2026, 3, 10)


@pytest.mark.parametrize(
    "url", ["http://127.0.0.1/admin", "http://169.254.169.254/latest", "file:///etc/passwd"]
)
def test_destino_interno_e_recusado(url):
    from apps.knowledge.web import PaginaIndisponivel, conferir_destino

    with pytest.raises(PaginaIndisponivel):
        conferir_destino(url)


@pytest.mark.django_db
def test_ingerir_url_guarda_a_copia_e_deduplica(tenant_de_fontes, monkeypatch):
    from apps.knowledge.entradas import ingerir_url
    from apps.knowledge.flows import passo_converter

    chamadas = []

    def baixar(url):
        chamadas.append(url)
        return PAGINA.encode(), url, "text/html"

    monkeypatch.setattr("apps.knowledge.entradas.baixar", baixar)
    categoria = _categoria("veiculo")
    url = "https://blog.exemplo.com.br/bdi"

    primeiro = ingerir_url(url, category=categoria, iniciar=False)
    segundo = ingerir_url(url, category=categoria, iniciar=False)

    assert not primeiro.ja_existia and segundo.ja_existia
    assert len(chamadas) == 1
    documento = primeiro.document
    assert documento.origin == Document.Origin.URL
    assert documento.fetched_at is not None

    from types import SimpleNamespace

    passo_converter(SimpleNamespace(target_object_id=str(documento.pk)))
    documento.refresh_from_db()
    assert documento.status == Document.Status.PENDING_CURATION
    assert documento.extraction_method == "web"
    assert documento.published_on == datetime.date(2026, 3, 10)
    assert documento.title == "Como calcular o BDI"


# ---------------------------------------------------------------------------
# Nota do especialista
# ---------------------------------------------------------------------------
@pytest.mark.django_db
def test_nota_do_especialista_ja_entra_curada_e_atribuida(ambiente):  # noqa: F811
    _, _, client = ambiente

    resposta = client.post(
        reverse("knowledge:nota", urlconf="core.urls_tenants"),
        {
            "titulo": "Quando refazer o orcamento",
            "autor": "Eng. Paulo Lima",
            "credencial": "engenheiro civil",
            "texto": "## Mudanca de projeto\n\nToda mudanca de projeto pede orcamento novo.",
        },
    )

    assert resposta.status_code == 302
    documento = Document.objects.get(origin=Document.Origin.NOTE)
    assert documento.status == Document.Status.CURATED
    assert documento.category.source_class == "especialista"
    trecho = documento.chunks.get()
    assert trecho.source_label == "Eng. Paulo Lima, engenheiro civil"
    assert trecho.citation_mode == "atribuicao"


# ---------------------------------------------------------------------------
# Curadoria e categorias
# ---------------------------------------------------------------------------
def _form_de_curadoria(categoria, **dados):
    from apps.knowledge.forms import CuradoriaDeDocumento

    documento = Document(category=categoria)
    base = {"title": "Manual", "license": "unknown", "authority_score": 50}
    return CuradoriaDeDocumento({**base, **dados}, instance=documento)


@pytest.mark.django_db
def test_material_da_empresa_nao_exige_url_mas_exige_um_nome(tenant_de_fontes):
    categoria = _categoria("empresa")

    sem_nome = _form_de_curadoria(categoria)
    com_rotulo = _form_de_curadoria(categoria, source_label="Manual de uso, v3")

    assert not sem_nome.is_valid()
    assert "source_label" in sem_nome.errors
    assert com_rotulo.is_valid(), com_rotulo.errors


@pytest.mark.django_db
def test_artigo_cientifico_continua_exigindo_url_e_autores(tenant_de_fontes):
    form = _form_de_curadoria(_categoria("cientifico"))

    assert not form.is_valid()
    assert {"source_url", "authors"} <= set(form.errors)


@pytest.mark.django_db
def test_categorias_padrao_e_edicao_pela_tela(ambiente):  # noqa: F811
    _, _, client = ambiente
    url = reverse("knowledge:categorias", urlconf="core.urls_tenants")

    client.post(url, {"acao": "padrao"})
    comunidade = DocumentCategory.objects.get(source_class="comunidade")
    assert not comunidade.supports_central_idea

    client.post(
        f"{url}?editar={comunidade.pk}",
        {
            "name": "Forum de engenharia",
            "source_class": "comunidade",
            "citation_mode": "interna",
            "validity_days": "180",
            "description": "",
        },
    )
    comunidade.refresh_from_db()
    assert comunidade.name == "Forum de engenharia"
    assert comunidade.validity_days == 180
    assert client.get(url).status_code == 200
